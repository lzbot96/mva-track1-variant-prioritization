"""Extract paragraphs/tables locally; no clinical content emitted to stdout."""
import json,re
from docx import Document
from project_paths import project_root,required_env
R=project_root();src=R/required_env('MVA_CLINICAL_DOCX');doc=Document(src)
paragraphs=[{'paragraph':i+1,'text':p.text} for i,p in enumerate(doc.paragraphs) if p.text.strip()]
tables=[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
text='\n'.join([x['text'] for x in paragraphs]+[c for t in tables for row in t for c in row])
hpo=sorted(set(re.findall(r'HP:\d{7}',text)))
(R/'results/phenotype/clinical_document_extracted.json').write_text(json.dumps({'source':src.name,'paragraphs':paragraphs,'tables':tables,'hpo_ids_explicit':hpo},indent=2)+'\n')
print('Clinical document extracted locally; content not emitted.')
