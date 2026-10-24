"""Synthetic document only; needs python-docx. No real clinical content."""
import json,os,subprocess,sys,tempfile
from pathlib import Path
from docx import Document
R=Path(__file__).resolve().parents[1]
with tempfile.TemporaryDirectory(prefix='mva-clinical-test-') as td:
    t=Path(td);(t/'results/phenotype').mkdir(parents=True)
    terms=['HP:'+str(n).zfill(7) for n in [1,2]]
    doc=Document();doc.add_paragraph('Synthetic fixture')
    table=doc.add_table(rows=1,cols=4)
    for c,s in zip(table.rows[0].cells,['feature','term','hpo','note']):c.text=s
    for i,term in enumerate(terms):
        for c,s in zip(table.add_row().cells,['synthetic','synthetic',term,'fixture']):c.text=s
    doc.save(t/'synthetic.docx')
    env={**os.environ,'PYTHONDONTWRITEBYTECODE':'1','MVA_PROJECT_ROOT':td,'MVA_CLINICAL_DOCX':'synthetic.docx','MVA_FAMILY_HISTORY_HPO':terms[1]}
    for script in ['extract_clinical.py','curate_hpo.py']:
        p=subprocess.run([sys.executable,str(R/'scripts'/script)],env=env,capture_output=True,text=True)
        assert p.returncode==0,p.stderr
        assert all(term not in p.stdout for term in terms)
    extracted=json.loads((t/'results/phenotype/clinical_document_extracted.json').read_text())
    curated=json.loads((t/'results/phenotype/curated_hpo.json').read_text())
    assert extracted['hpo_ids_explicit']==terms
    assert curated['all_challenge_hpo']==terms and curated['proband_only_hpo']==terms[:1]
    assert curated['explicit_absent_hpo']==[]
print('PASS: synthetic DOCX table extraction, explicit IDs, reviewer family assignment, no sensitive stdout.')
